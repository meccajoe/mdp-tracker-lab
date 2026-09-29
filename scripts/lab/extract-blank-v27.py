"""Read the exported blank v27 template; never contact or modify Google Sheets.

Usage: python scripts/lab/extract-blank-v27.py /path/to/export.xlsx
Writes the independently sourced, zero-estimate new-quote template.
"""
import hashlib
import json
import pathlib
import sys

import openpyxl

source = pathlib.Path(sys.argv[1])
raw = openpyxl.load_workbook(source, data_only=False)
cached = openpyxl.load_workbook(source, data_only=True)
settings_rows = {
    'contingency': 4, 'opex': 5, 'indirect': 6, 'laborSell': 8,
    'materialMarkup': 11, 'shopDay': 12, 'efficiency': 13,
    'graphicsSell': 16, 'graphicsCost': 17, 'handlingMinutes': 18,
    'handlingCrew': 19, 'designSell': 22, 'pmFee': 23, 'pmBonus': 25,
    'leadDay': 27, 'supportDay': 28, 'travelFactor': 29, 'pmTravelDay': 30,
    'siteHours': 31, 'supportCost': 32, 'offFactor': 33,
    'equipmentMarkup': 36, 'resaleMarkup': 37, 'travelMarkup': 38,
    'freightMarkup': 39, 'burdenMultiplier': 64,
}
settings = {key: cached['Settings'][f'B{row}'].value for key, row in settings_rows.items()}
settings['burdenedRateOverride'] = None if raw['Settings']['B7'].data_type == 'f' else cached['Settings']['B7'].value
trades = [{'id': f'trade-{r}', 'name': cached['Settings'][f'A{r}'].value,
           'wage': cached['Settings'][f'B{r}'].value}
          for r in range(66, 78) if cached['Settings'][f'A{r}'].value]
catalog = []
for r in list(range(5, 1005)) + list(range(1009, 1059)):
    name, unit, cost = [cached['Materials DB'].cell(r, c).value for c in (1, 2, 3)]
    if not isinstance(name, str) or not name.strip():
        continue
    if cost is not None and not isinstance(cost, (int, float)):
        raise ValueError(f'Unresolved catalog cost: row {r}')
    catalog.append({'id': f'material-{r}', 'name': name, 'unit': unit or '', 'unitCost': cost or 0})
lines = []
for r in range(5, 42):
    name, kind = [cached['Quote Builder'].cell(r, c).value for c in (1, 2)]
    lines.append({'id': f'line-{r}', 'name': name or f'Spare {r - 19}',
                  'type': kind or 'Fabrication', 'takeoffDriven': True,
                  'inputs': {key: 0 for key in ('materials','resale','hours','days','sqft','panels','rental','siteDays','travelDays','cost')},
                  'overrides': {}, 'priceOverride': None})
digest = hashlib.sha256(source.read_bytes()).hexdigest()
document = {'schemaVersion': 1, 'assumptionsVersion': 'blank-v27-' + digest[:12],
            'commission': 0, 'settings': settings, 'trades': trades, 'catalog': catalog,
            'lines': lines, 'takeoffs': []}
output = pathlib.Path('src/data/quote-v27-blank.json')
output.write_text(json.dumps(document, indent=2) + '\n')
print(f'Extracted {len(lines)} lines and {len(catalog)} catalog entries; source SHA256 {digest}')
